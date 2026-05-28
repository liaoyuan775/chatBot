from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


def set_run_fonts(run, east_asia: str = "宋体", latin: str = "Times New Roman", size: int = 12):
    run.font.name = latin
    run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def insert_paragraph_before(paragraph, text: str, style_name: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style_name
    run = new_para.add_run(text)
    set_run_fonts(run)
    return new_para


def remove_paragraph(paragraph):
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


base = Path(r"D:\AgentLearning\chatBot\output\doc")
docx_path = max(base.glob("*.docx"), key=lambda p: p.stat().st_size)
doc = Document(str(docx_path))

# 目录为静态条目，正文从该一级标题开始。
body_start_idx = next(
    i for i, p in enumerate(doc.paragraphs) if p.style.name == "Heading 1" and p.text.strip() == "1 绪论"
)
anchor = doc.paragraphs[body_start_idx]

# 删除目录末尾结论之后到正文开始之前的空白段落，避免目录页尾部留出无意义空行。
for idx in range(body_start_idx - 1, -1, -1):
    p = doc.paragraphs[idx]
    if p.text.strip() == "":
        remove_paragraph(p)
        continue
    break

# 如果目录中还没有补齐后置章节，则手工补齐。
existing_toc_texts = {p.text.strip() for p in doc.paragraphs[:body_start_idx]}
appendix_entries = [
    "参考文献\t39",
    "致谢\t41",
    "附录1 部分主要源代码\t43",
]
for entry in appendix_entries:
    if entry not in existing_toc_texts:
        insert_paragraph_before(anchor, entry, "toc 1")

doc.save(str(docx_path))
print(docx_path)
