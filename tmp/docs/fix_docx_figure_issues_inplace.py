from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image


ROOT = Path(r"D:/AgentLearning/chatBot")
DOC_DIR = ROOT / "output" / "doc"
DIAGRAM_DIR = DOC_DIR / "thesis_diagrams"
ASSET_DIR = DOC_DIR / "thesis_assets"
TMP_DIR = ROOT / "tmp" / "docs"


def find_main_docx() -> Path:
    candidates = [
        p
        for p in DOC_DIR.glob("*.docx")
        if "附件" not in p.name and p.stat().st_size > 100000
    ]
    if not candidates:
        raise FileNotFoundError("main thesis docx not found")
    return max(candidates, key=lambda p: p.stat().st_size)


def get_paragraphs(doc: Document):
    return doc.paragraphs


def find_paragraph(doc: Document, exact: str):
    for p in get_paragraphs(doc):
        if p.text.strip() == exact:
            return p
    raise ValueError(f"paragraph not found: {exact}")


def clear_paragraph(paragraph):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def set_paragraph_text(paragraph, text: str):
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def paragraph_before(doc: Document, paragraph):
    paragraphs = get_paragraphs(doc)
    for i, p in enumerate(paragraphs):
        if p._p is paragraph._p:
            if i == 0:
                raise ValueError("paragraph has no previous sibling")
            return paragraphs[i - 1]
    raise ValueError("paragraph not found in document")


def move_block_after(anchor, block):
    anchor_el = anchor._p
    for paragraph in reversed(block):
        el = paragraph._p
        el.getparent().remove(el)
        anchor_el.addnext(el)


def draw_mapping(doc: Document):
    draw_paragraphs = [p for p in get_paragraphs(doc) if p._p.xpath(".//w:drawing")]
    return {p._p: shape for p, shape in zip(draw_paragraphs, doc.inline_shapes)}


def set_shape_size_for_paragraph(shape_map, paragraph, width_in: float, image_path: Path | None = None):
    shape = shape_map[paragraph._p]
    if image_path is not None:
        img = Image.open(image_path)
        ratio = img.size[1] / img.size[0]
    else:
        ratio = shape.height / shape.width
    shape.width = Inches(width_in)
    shape.height = int(shape.width * ratio)


def insert_picture(paragraph, image_path: Path, width_in: float):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_in))


def backup_file(path: Path) -> Path:
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    backup = TMP_DIR / f"{path.stem}.before_figure_fix{path.suffix}"
    shutil.copy2(path, backup)
    return backup


def main():
    docx_path = find_main_docx()
    backup_path = backup_file(docx_path)
    doc = Document(str(docx_path))

    # Existing drawings before we move anything.
    shape_map = draw_mapping(doc)

    # Enlarge 3-1 and 4-1 without touching surrounding styles.
    p_3_1_caption = find_paragraph(doc, "图3-1 数据库 E-R 关系图")
    p_3_1_image = paragraph_before(doc, p_3_1_caption)
    set_shape_size_for_paragraph(shape_map, p_3_1_image, 6.1, DIAGRAM_DIR / "database_er.png")

    p_4_1_caption = find_paragraph(doc, "图4-1 系统总体架构图")
    p_4_1_image = paragraph_before(doc, p_4_1_caption)
    set_shape_size_for_paragraph(shape_map, p_4_1_image, 6.25, DIAGRAM_DIR / "system_architecture.png")

    # Move the misplaced chain-config block so figure order in chapter 5 is sequential.
    p_chain_text = find_paragraph(doc, "链路配置界面用于把一体化链路与可拆分链路绑定到具体模型、人格和音色，相关页面如图5-4所示。")
    p_chain_image = paragraph_before(doc, find_paragraph(doc, "图5-1 链路配置页面"))
    p_chain_caption = find_paragraph(doc, "图5-1 链路配置页面")
    set_paragraph_text(p_chain_caption, "图5-4 链路配置页面")

    p_persona_caption = find_paragraph(doc, "图5-3 人格配置页面")
    move_block_after(p_persona_caption, [p_chain_text, p_chain_image, p_chain_caption])

    # Rebuild mapping after moving existing drawings.
    shape_map = draw_mapping(doc)

    # 4-3 image is missing in the current docx. Insert it into the existing blank paragraph.
    p_4_3_caption = find_paragraph(doc, "图4-3 双语音链路处理流程图")
    p_4_3_image = paragraph_before(doc, p_4_3_caption)
    insert_picture(p_4_3_image, DIAGRAM_DIR / "voice_interaction_flow.png", 4.55)

    # Keep 5-1 readable.
    p_5_1_caption = find_paragraph(doc, "图5-1 实时语音通话时序图")
    p_5_1_image = paragraph_before(doc, p_5_1_caption)
    set_shape_size_for_paragraph(shape_map, p_5_1_image, 6.0)

    # Move system-settings block before section 5.10 so 5-9/5-10 stay in order.
    p_settings_text = find_paragraph(doc, "系统配置页面集中展示数据库状态、异常策略和高危清理操作，其界面如图5-9所示。")
    p_settings_image = paragraph_before(doc, find_paragraph(doc, "图5-9 系统配置页面"))
    p_settings_caption = find_paragraph(doc, "图5-9 系统配置页面")
    p_realtime_heading = find_paragraph(doc, "5.10 实时 WebSocket 服务设计")
    move_block_after(paragraph_before(doc, p_realtime_heading), [p_settings_text, p_settings_image, p_settings_caption])

    # 5-10 image is also missing. Fill the blank paragraph before its caption.
    p_5_10_caption = find_paragraph(doc, "图5-10 实时语音对话与打断交互示意图")
    p_5_10_image = paragraph_before(doc, p_5_10_caption)
    insert_picture(p_5_10_image, ASSET_DIR / "fig5-10-realtime-interrupt-scene.png", 6.2)

    doc.save(str(docx_path))

    print(f"UPDATED: {docx_path}")
    print(f"BACKUP: {backup_path}")


if __name__ == "__main__":
    main()
