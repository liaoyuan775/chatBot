from __future__ import annotations

import json
import re
import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "doc"
SRC_DOCX = OUT / "智能语音聊天机器人设计与实现_本科毕业论文.docx"
FINAL_DOCX = OUT / "智能语音聊天机器人设计与实现.docx"
SRC_ATTACHMENT = OUT / "智能语音聊天机器人设计与实现_图表源码附件.docx"
FINAL_ATTACHMENT = OUT / "智能语音聊天机器人设计与实现-附件.docx"
FINAL_MD = OUT / "智能语音聊天机器人设计与实现.md"
CHECK_JSON = OUT / "智能语音聊天机器人设计与实现-交付自检.json"


TITLE = "智能语音聊天机器人设计与实现"
TODAY_CN = "2026年5月1日"


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_paragraph_before(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style:
        new_para.style = style
    return new_para


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style:
        new_para.style = style
    return new_para


def set_run_font(run, east_asia: str = "宋体", ascii_font: str = "Times New Roman", size: int = 12, bold: bool = False) -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold


def format_body(paragraph: Paragraph, first_indent: bool = True) -> None:
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24) if first_indent else Pt(0)
    for run in paragraph.runs:
        set_run_font(run, "宋体", "Times New Roman", 12)


def format_center_title(paragraph: Paragraph, size: int = 18) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(22)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        set_run_font(run, "黑体", "Times New Roman", size, True)


def format_heading(paragraph: Paragraph, level: int) -> None:
    paragraph.style = f"Heading {level}"
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.first_line_indent = Pt(0)
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        size = 18
    elif level == 2:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        size = 15
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        size = 14
    for run in paragraph.runs:
        set_run_font(run, "黑体", "Times New Roman", size, True)


def add_page_break(paragraph: Paragraph) -> None:
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_front_matter(doc: Document) -> None:
    first = doc.paragraphs[0]
    cover_lines = [
        ("安徽工业大学", 22),
        ("毕业设计（论文）说明书", 22),
        ("", 12),
        (f"题目：{TITLE}", 16),
        ("专业：计算机科学与技术", 14),
        ("班级：________________", 14),
        ("姓名：________________", 14),
        ("学号：________________", 14),
        ("指导教师：________________", 14),
        (f"日期：{TODAY_CN}", 14),
    ]
    anchor = first
    for text, size in cover_lines:
        p = insert_paragraph_before(anchor, text)
        if text:
            format_center_title(p, size)
        else:
            format_body(p, first_indent=False)
    p = insert_paragraph_before(anchor, "")
    add_page_break(p)

    task_blocks = [
        ("毕业设计（论文）任务书", "title"),
        (f"题目：{TITLE}", "body"),
        ("主要任务：围绕智能语音聊天机器人系统完成需求分析、数据库设计、总体架构设计、详细模块设计、编码实现与测试验证。论文重点说明统一会话运行时、分阶段语音链路、一体化实时语音链路、RAG 知识增强、人格配置、音色配置和实时 WebSocket 通话等内容。", "body"),
        ("技术路线：前端采用 React、Vite 与 TypeScript 构建交互界面；后端采用 FastAPI 组织会话、消息、知识检索、模型路由和实时通信服务；数据层采用 PostgreSQL 与 pgvector 保存会话数据、记忆数据、知识库向量、链路配置、音色配置和运行审计数据。", "body"),
        ("完成要求：系统应能够支持文本对话、语音输入、实时语音通话、知识库问答、上下文记忆、Provider 与模型配置、链路配置、人格配置、音色管理、图片生成和系统设置等能力；论文应包含 E-R 图、系统架构图、流程图、真实界面截图、测试用例和参考文献。", "body"),
        ("进度安排：完成项目需求分析与资料整理；完成系统设计、数据库设计与核心模块实现；完成系统测试、论文正文撰写、图表整理、格式检查和初稿提交。", "body"),
        ("主要参考资料：学校论文书写要求与说明书模板、项目源码、项目截图、PostgreSQL 与 pgvector 文档、FastAPI 文档、React 文档以及语音识别、检索增强生成和长时记忆相关研究文献。", "body"),
    ]
    for text, kind in task_blocks:
        p = insert_paragraph_before(anchor, text)
        if kind == "title":
            format_center_title(p, 18)
        else:
            format_body(p)
    p = insert_paragraph_before(anchor, "")
    add_page_break(p)


def strip_old_front_matter(doc: Document) -> None:
    # Keep everything from the first declaration onward; the old cover omitted
    # several required school-template fields, so it is rebuilt above.
    idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "毕业设计（论文）独创性声明":
            idx = i
            break
    if idx is None:
        return
    for p in list(doc.paragraphs[:idx]):
        remove_paragraph(p)


def normalize_declaration_dates(doc: Document) -> None:
    for p in doc.paragraphs:
        text = p.text.strip()
        if re.fullmatch(r"2026年\s+\d+\s+月\s+\d+\s+日", text):
            p.clear()
            p.add_run(TODAY_CN)
            format_body(p, first_indent=False)


def rebuild_toc(doc: Document) -> None:
    toc_idx = None
    first_chapter_idx = None
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text == "目 录":
            toc_idx = i
        if text == "1 绪论":
            first_chapter_idx = i
            break
    if toc_idx is None or first_chapter_idx is None:
        return

    old_toc = list(doc.paragraphs[toc_idx + 1 : first_chapter_idx])
    for p in old_toc:
        remove_paragraph(p)

    toc_title = doc.paragraphs[toc_idx]
    format_center_title(toc_title, 18)
    toc_field = insert_paragraph_after(toc_title)
    toc_field.paragraph_format.first_line_indent = Pt(0)
    toc_field.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "目录将在 Word 中自动更新"
    run_text = OxmlElement("w:r")
    run_text.append(text)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for node in (fld_begin, instr, fld_sep, run_text, fld_end):
        r = OxmlElement("w:r")
        if node.tag.endswith("}r"):
            toc_field._p.append(node)
        else:
            r.append(node)
            toc_field._p.append(r)
    p = insert_paragraph_after(toc_field)
    add_page_break(p)


def normalize_styles(doc: Document) -> None:
    chapter_titles = {
        "1 绪论",
        "2 开发环境及相关技术介绍",
        "3 系统需求分析与数据库设计",
        "4 系统总体设计",
        "5 系统详细设计",
        "6 系统测试",
        "7 结论",
    }
    center_titles = {
        "摘  要",
        "Abstract",
        "目 录",
        "参考文献",
        "致谢",
        "附录",
        "英文资料及翻译",
        "毕业设计（论文）独创性声明",
    }
    for section in doc.sections:
        section.start_type = WD_SECTION.NEW_PAGE
        section.top_margin = Cm(2.8)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        header = section.header.paragraphs[0]
        header.text = "安徽工业大学毕业设计（论文）"
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in header.runs:
            set_run_font(run, "宋体", "Times New Roman", 14)

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if text in chapter_titles:
            format_heading(p, 1)
        elif p.style and p.style.name == "Heading 2":
            format_heading(p, 2)
        elif p.style and p.style.name == "Heading 3":
            if text:
                format_heading(p, 3)
        elif text in center_titles:
            format_center_title(p, 18)
        elif text.startswith("图"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            for run in p.runs:
                set_run_font(run, "宋体", "Times New Roman", 12)
        elif text.startswith("表"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            for run in p.runs:
                set_run_font(run, "宋体", "Times New Roman", 12)
        elif not (p.style and p.style.name.lower().startswith("toc")):
            format_body(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    p.paragraph_format.line_spacing = Pt(18)
                    p.paragraph_format.first_line_indent = Pt(0)
                    for run in p.runs:
                        set_run_font(run, "宋体", "Times New Roman", 10)


def remove_empty_heading(doc: Document) -> None:
    for p in list(doc.paragraphs):
        if p.style and p.style.name.startswith("Heading") and not p.text.strip():
            remove_paragraph(p)


def copy_derivatives() -> None:
    if SRC_ATTACHMENT.exists():
        shutil.copy2(SRC_ATTACHMENT, FINAL_ATTACHMENT)
    old_md = OUT / "智能语音聊天机器人设计与实现_本科毕业论文.md"
    if old_md.exists():
        shutil.copy2(old_md, FINAL_MD)

    old_image_map = OUT / "智能语音聊天机器人设计与实现_本科毕业论文-image-map.json"
    if old_image_map.exists():
        shutil.copy2(old_image_map, OUT / "智能语音聊天机器人设计与实现-image-map.json")
    old_refs = OUT / "智能语音聊天机器人设计与实现_本科毕业论文-文献核验清单.json"
    if old_refs.exists():
        shutil.copy2(old_refs, OUT / "智能语音聊天机器人设计与实现-文献核验清单.json")


def validate(doc: Document) -> dict:
    text = "\n".join(p.text for p in doc.paragraphs)
    h1 = [p.text.strip() for p in doc.paragraphs if p.style and p.style.name == "Heading 1" and p.text.strip()]
    h2_count = sum(1 for p in doc.paragraphs if p.style and p.style.name == "Heading 2" and p.text.strip())
    h3_count = sum(1 for p in doc.paragraphs if p.style and p.style.name == "Heading 3" and p.text.strip())
    sensitive = ["API_KEY", "sk-", ".env", "POSTGRES_DSN", "VOLCENGINE", "DASHSCOPE_API_KEY", "SILICONFLOW_API_KEY", "175.27."]
    placeholders = ["XXX", "YYY", "ZZZ", "N种权限", "定稿前删除"]
    return {
        "title": TITLE,
        "main_docx": str(FINAL_DOCX),
        "attachment_docx": str(FINAL_ATTACHMENT),
        "paragraphs": len([p for p in doc.paragraphs if p.text.strip()]),
        "tables": len(doc.tables),
        "images": sum(1 for rel in doc.part.rels.values() if "image" in rel.reltype),
        "non_space_chars": len(re.sub(r"\s+", "", text)),
        "heading_1": h1,
        "heading_2_count": h2_count,
        "heading_3_count": h3_count,
        "has_task_book": "毕业设计（论文）任务书" in text,
        "has_toc": "目 录" in text,
        "has_er_diagram_caption": "数据库 E-R" in text,
        "sensitive_hits": {k: text.count(k) for k in sensitive if text.count(k)},
        "placeholder_hits": {k: text.count(k) for k in placeholders if text.count(k)},
        "final_date": TODAY_CN,
    }


def main() -> None:
    if not SRC_DOCX.exists():
        raise FileNotFoundError(SRC_DOCX)
    shutil.copy2(SRC_DOCX, FINAL_DOCX)
    doc = Document(FINAL_DOCX)
    remove_empty_heading(doc)
    strip_old_front_matter(doc)
    add_front_matter(doc)
    normalize_declaration_dates(doc)
    rebuild_toc(doc)
    normalize_styles(doc)
    doc.save(FINAL_DOCX)
    copy_derivatives()
    result = validate(Document(FINAL_DOCX))
    CHECK_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
