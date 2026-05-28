from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:/AgentLearning/chatBot")
TMP_DIR = ROOT / "tmp" / "docs"
OUT_DIR = ROOT / "output" / "doc"
DIAGRAM_DIR = OUT_DIR / "thesis_diagrams"
ASSET_DIR = OUT_DIR / "thesis_assets"

TITLE = "智能语音聊天机器人设计与实现"
COLLEGE = "计算机科学与技术学院"
MAJOR = "计算机科学与技术"
CLASS_NAME = "________________"
AUTHOR = "廖园"
STUDENT_ID = "229074182"
SUPERVISOR = "王朋飞"
TODAY_CN = "2026年5月16日"
HEADER_TEXT = "安徽工业大学毕业设计（论文）说明书"

MAIN_MD = OUT_DIR / f"{TITLE}.md"
MAIN_DOCX = OUT_DIR / f"{TITLE}.docx"
ATTACHMENT_DOCX = OUT_DIR / f"{TITLE}-附件.docx"
IMAGE_MAP_JSON = OUT_DIR / f"{TITLE}-image-map.json"
REFS_JSON = OUT_DIR / f"{TITLE}-文献核验清单.json"
CHECK_JSON = OUT_DIR / f"{TITLE}-交付自检.json"


def locate_source_docx() -> Path:
    candidates = list(TMP_DIR.glob("*本科毕业论文.docx"))
    if not candidates:
        candidates = list(TMP_DIR.glob("*.docx"))
    if not candidates:
        raise FileNotFoundError("未找到可用的论文源稿 docx")
    return max(candidates, key=lambda p: p.stat().st_size)


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_paragraph_before(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def clear_paragraph(paragraph: Paragraph) -> None:
    paragraph.clear()
    paragraph._p[:] = []


def set_run_font(
    run,
    east_asia: str = "宋体",
    ascii_font: str = "Times New Roman",
    size: float = 12,
    bold: bool = False,
) -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold


def apply_run_fonts(paragraph: Paragraph, east_asia: str, ascii_font: str, size: float, bold: bool = False) -> None:
    if not paragraph.runs:
        paragraph.add_run("")
    for run in paragraph.runs:
        set_run_font(run, east_asia=east_asia, ascii_font=ascii_font, size=size, bold=bold)


def format_body(paragraph: Paragraph, first_indent: bool = True) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24) if first_indent else Pt(0)
    apply_run_fonts(paragraph, "宋体", "Times New Roman", 12, False)


def format_cover_line(paragraph: Paragraph, size: float, bold: bool = False) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(0)
    apply_run_fonts(paragraph, "黑体" if bold else "宋体", "Times New Roman", size, bold)


def format_center_title(paragraph: Paragraph, size: float = 18) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.first_line_indent = Pt(0)
    apply_run_fonts(paragraph, "黑体", "Times New Roman", size, True)


def format_heading(paragraph: Paragraph, level: int) -> None:
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.first_line_indent = Pt(0)
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.style = "Heading 1"
        apply_run_fonts(paragraph, "黑体", "Times New Roman", 16, True)
    elif level == 2:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.style = "Heading 2"
        apply_run_fonts(paragraph, "黑体", "Times New Roman", 14, True)
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.style = "Heading 3"
        apply_run_fonts(paragraph, "黑体", "Times New Roman", 12, True)


def format_caption(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(0)
    apply_run_fonts(paragraph, "宋体", "Times New Roman", 10.5, False)


def format_code_block(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(14)
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    pf.first_line_indent = Pt(0)
    apply_run_fonts(paragraph, "等线", "Consolas", 9.5, False)


def add_page_break(paragraph: Paragraph) -> None:
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_page_number(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)
    set_run_font(run, "宋体", "Times New Roman", 10.5, False)


def is_code_like(text: str) -> bool:
    text = text.strip()
    if not text:
        return False
    signals = ["=>", "await ", "const ", "def ", "async ", "select(", "if (", "{", "}", "return ", "db.add("]
    hit_count = sum(1 for token in signals if token in text)
    return "\n" in text or hit_count >= 2


def strip_old_front_matter(doc: Document) -> Paragraph:
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip() == "毕业设计（论文）独创性声明":
            anchor = p
            break
    if anchor is None:
        raise ValueError("未找到独创性声明页，无法重建前置页")
    for p in list(doc.paragraphs):
        if p is anchor:
            break
        remove_paragraph(p)
    return anchor


def insert_cover_and_taskbook(anchor: Paragraph) -> None:
    cover_lines = [
        ("安徽工业大学", 22, True),
        ("毕业设计（论文）说明书", 22, True),
        ("", 12, False),
        (TITLE, 18, True),
        ("", 12, False),
        (f"学院：{COLLEGE}", 14, False),
        (f"专业：{MAJOR}", 14, False),
        (f"班级：{CLASS_NAME}", 14, False),
        (f"姓名：{AUTHOR}", 14, False),
        (f"学号：{STUDENT_ID}", 14, False),
        (f"指导教师：{SUPERVISOR}", 14, False),
        (f"日期：{TODAY_CN}", 14, False),
    ]
    for text, size, bold in reversed(cover_lines):
        p = insert_paragraph_before(anchor, text)
        format_cover_line(p, size, bold) if text else format_body(p, first_indent=False)
    p = insert_paragraph_before(anchor, "")
    add_page_break(p)

    task_lines = [
        ("毕业设计（论文）任务书", "title"),
        (f"题目：{TITLE}", "body"),
        ("主要任务：围绕智能语音聊天机器人系统完成需求分析、数据库设计、总体架构设计、详细模块设计、功能实现与系统测试，形成符合学校模板要求的毕业设计论文。", "body"),
        ("研究重点：重点说明统一会话运行时 SessionRuntime、integrated-realtime 与 split-chain 双语音策略、WebSocket 实时通话、RAG 检索增强、上下文记忆、Provider 与模型管理、人格配置和音色管理等关键设计。", "body"),
        ("技术路线：前端采用 React、Vite 与 TypeScript 构建界面与运行时；后端采用 FastAPI、异步 SQLAlchemy、PostgreSQL 与 pgvector 构建会话服务、知识增强和实时通信能力。", "body"),
        ("完成要求：论文需包含中英文摘要、目录、需求分析、数据库设计、总体设计、详细设计与实现、系统测试、参考文献、致谢，并配套系统架构图、功能模块图、语音流程图、实时通话时序图、数据库 E-R 图和真实页面截图。", "body"),
    ]
    for text, kind in reversed(task_lines):
        p = insert_paragraph_before(anchor, text)
        if kind == "title":
            format_center_title(p, 18)
        else:
            format_body(p)
    p = insert_paragraph_before(anchor, "")
    add_page_break(p)


def normalize_dates(doc: Document) -> None:
    for p in doc.paragraphs:
        text = p.text.strip()
        if re.fullmatch(r"2026年\s*\d+\s*月\s*\d+\s*日", text) or re.fullmatch(r"二O二六\s*年\s*五\s*月\s*二十七\s*日", text):
            clear_paragraph(p)
            p.add_run(TODAY_CN)
            format_body(p, first_indent=False)


def rename_chapter_titles(doc: Document) -> None:
    mapping = {
        "1 绪论": "第1章 绪论",
        "2 开发环境及相关技术介绍": "第2章 开发环境及相关技术介绍",
        "3 系统需求分析与数据库设计": "第3章 系统需求分析与数据库设计",
        "4 系统总体设计": "第4章 系统总体设计",
        "5 系统详细设计": "第5章 系统详细设计与实现",
        "6 系统测试": "第6章 系统测试",
        "7 结论": "结 论",
    }
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in mapping:
            clear_paragraph(p)
            p.add_run(mapping[text])


def rebuild_toc(doc: Document) -> None:
    toc_title = None
    chapter_start = None
    for p in doc.paragraphs:
        text = p.text.strip()
        if text == "目 录":
            toc_title = p
        elif text == "第1章 绪论":
            chapter_start = p
            break
    if toc_title is None or chapter_start is None:
        return

    in_between = []
    passed = False
    for p in doc.paragraphs:
        if p is toc_title:
            passed = True
            continue
        if not passed:
            continue
        if p is chapter_start:
            break
        in_between.append(p)
    for p in in_between:
        remove_paragraph(p)

    format_center_title(toc_title, 18)
    toc_field = insert_paragraph_after(toc_title)
    toc_field.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc_field.paragraph_format.first_line_indent = Pt(0)
    toc_field.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    toc_field.paragraph_format.line_spacing = Pt(18)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder_run = toc_field.add_run("目录将在 Word 中自动更新")
    set_run_font(placeholder_run, "宋体", "Times New Roman", 12, False)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    toc_field._p.insert(0, OxmlElement("w:r"))
    toc_field._p[0].append(fld_begin)
    toc_field._p.insert(1, OxmlElement("w:r"))
    toc_field._p[1].append(instr)
    toc_field._p.insert(2, OxmlElement("w:r"))
    toc_field._p[2].append(fld_sep)
    toc_field._p.append(OxmlElement("w:r"))
    toc_field._p[-1].append(fld_end)
    page_break = insert_paragraph_after(toc_field, "")
    add_page_break(page_break)


def normalize_document_styles(doc: Document) -> None:
    chapter_titles = {
        "第1章 绪论",
        "第2章 开发环境及相关技术介绍",
        "第3章 系统需求分析与数据库设计",
        "第4章 系统总体设计",
        "第5章 系统详细设计与实现",
        "第6章 系统测试",
        "结 论",
    }
    center_titles = {
        "摘  要",
        "Abstract",
        "目 录",
        "参考文献",
        "致谢",
        "毕业设计（论文）独创性声明",
        "毕业设计（论文）任务书",
    }

    for section in doc.sections:
        section.start_type = WD_SECTION.NEW_PAGE
        section.top_margin = Cm(2.8)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        header = section.header.paragraphs[0]
        clear_paragraph(header)
        header.add_run(HEADER_TEXT)
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        apply_run_fonts(header, "宋体", "Times New Roman", 10.5, False)

        footer = section.footer.paragraphs[0]
        clear_paragraph(footer)
        add_page_number(footer)

    for p in doc.paragraphs:
        text = p.text.strip()
        style_name = p.style.name if p.style else ""
        if not text:
            continue
        if text in chapter_titles:
            format_heading(p, 1)
        elif re.match(r"^\d+\.\d+\s", text) or style_name == "Heading 2":
            format_heading(p, 2)
        elif re.match(r"^\d+\.\d+\.\d+\s", text) or style_name == "Heading 3":
            format_heading(p, 3)
        elif text in center_titles:
            format_center_title(p, 18)
        elif text.startswith("图"):
            format_caption(p)
        elif text.startswith("表"):
            format_caption(p)
        elif style_name.lower().startswith("toc"):
            continue
        elif is_code_like(text):
            format_code_block(p)
        else:
            format_body(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pf = p.paragraph_format
                    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    pf.line_spacing = Pt(18)
                    pf.first_line_indent = Pt(0)
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    apply_run_fonts(p, "宋体", "Times New Roman", 10.5, False)


def build_attachment_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.8)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    p = doc.add_paragraph()
    p.add_run(f"{TITLE} 附件")
    format_center_title(p, 18)

    intro = doc.add_paragraph()
    intro.add_run("本附件收录论文图表源码、关键数据库结构说明、主要接口说明以及图表来源说明，用于配合正文中的图文内容进行交叉核对。")
    format_body(intro)

    h1 = doc.add_paragraph()
    h1.add_run("1 图表 Mermaid 源码")
    format_heading(h1, 2)

    for idx, path in enumerate(sorted(DIAGRAM_DIR.glob("*.mmd")), start=1):
        h2 = doc.add_paragraph()
        h2.add_run(f"1.{idx} {path.name}")
        format_heading(h2, 3)
        code = doc.add_paragraph()
        code.add_run(path.read_text(encoding="utf-8"))
        format_code_block(code)

    h_db = doc.add_paragraph()
    h_db.add_run("2 关键数据库表结构说明")
    format_heading(h_db, 2)

    db_items = [
        "sessions：保存会话主记录，包括标题、当前链路、当前人格和基础时间戳。",
        "messages：保存用户消息、助手消息及其音频地址、消息状态和顺序关系。",
        "session_memories：保存短期原文记忆与中期摘要记忆，是上下文压缩的重要载体。",
        "long_term_memories：保存长期语义记忆向量，用于跨轮次、跨主题的相关事实回忆。",
        "provider_configs 与 model_catalog：分别保存厂商接入信息和同步后的模型目录，实现模型与业务调用解耦。",
        "knowledge_bases、knowledge_documents、knowledge_chunks：共同构成知识库实体、文档实体和切片向量实体。",
        "chains、voices、personas、call_configs：分别保存链路路由、音色、人格和调用参数配置。",
        "session_chain_audits：保存链路调用结果、回退原因和知识命中等审计信息。",
    ]
    for item in db_items:
        p = doc.add_paragraph()
        p.add_run(item)
        format_body(p)

    h_api = doc.add_paragraph()
    h_api.add_run("3 关键接口与实时通道说明")
    format_heading(h_api, 2)

    api_items = [
        "POST /api/chat/sessions：创建或查询会话。",
        "POST /api/chat/sessions/{session_id}/messages：提交文本消息并触发回复生成。",
        "GET /api/context/sessions/{session_id}/memory：查看当前会话记忆快照。",
        "POST /api/providers/{provider_name}/test：测试厂商连通性。",
        "POST /api/providers/{provider_name}/models/sync：同步远端模型目录。",
        "POST /api/knowledge/bases：创建知识库；POST /api/knowledge/documents：上传知识文档。",
        "POST /api/voices/{voice_id}/clone：执行音色克隆或样本处理。",
        "WS /ws/realtime-chat?session_id=...：建立实时语音会话，接收局部转写、文本增量和音频增量事件。",
    ]
    for item in api_items:
        p = doc.add_paragraph()
        p.add_run(item)
        format_body(p)

    h_src = doc.add_paragraph()
    h_src.add_run("4 图表来源说明")
    format_heading(h_src, 2)
    image_map = json.loads(IMAGE_MAP_JSON.read_text(encoding="utf-8"))
    for name, path in image_map.items():
        p = doc.add_paragraph()
        p.add_run(f"{name}：{path}")
        format_body(p)

    doc.save(ATTACHMENT_DOCX)


def write_reference_json() -> None:
    items = [
        {"index": 1, "title": "Attention Is All You Need", "type": "conference", "source": "https://papers.nips.cc/paper/7181-attention-is-all-you-need", "status": "selected"},
        {"index": 2, "title": "Robust Speech Recognition via Large-Scale Weak Supervision", "type": "preprint", "source": "https://arxiv.org/abs/2212.04356", "status": "selected"},
        {"index": 3, "title": "FunASR: A Fundamental End-to-End Speech Recognition Toolkit", "type": "conference", "source": "https://www.isca-archive.org/interspeech_2023/gao23b_interspeech.html", "status": "selected"},
        {"index": 4, "title": "Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks", "type": "conference", "source": "https://proceedings.neurips.cc/paper/2020/hash/6b493230205f780e1bc26945df7481e5-Abstract.html", "status": "selected"},
        {"index": 5, "title": "MemoryBank: Enhancing Large Language Models with Long-Term Memory", "type": "conference", "source": "https://ojs.aaai.org/index.php/AAAI/article/view/29946", "status": "selected"},
        {"index": 6, "title": "Retrieval-Augmented Generation for Large Language Models: A Survey", "type": "preprint", "source": "https://arxiv.org/abs/2312.10997", "status": "selected"},
        {"index": 7, "title": "FastAPI Documentation", "type": "official-doc", "source": "https://fastapi.tiangolo.com/", "status": "selected"},
        {"index": 8, "title": "React Documentation", "type": "official-doc", "source": "https://react.dev/", "status": "selected"},
        {"index": 9, "title": "PostgreSQL Documentation", "type": "official-doc", "source": "https://www.postgresql.org/docs/current/", "status": "selected"},
        {"index": 10, "title": "pgvector", "type": "official-doc", "source": "https://github.com/pgvector/pgvector", "status": "selected"},
        {"index": 11, "title": "The WebSocket Protocol (RFC 6455)", "type": "standard", "source": "https://datatracker.ietf.org/doc/html/rfc6455", "status": "selected"},
        {"index": 12, "title": "大语言模型的幻觉问题研究综述", "type": "journal", "source": "https://www.jos.org.cn/jos/article/abstract/pa049", "status": "selected"},
    ]
    payload = {
        "title": TITLE,
        "standard": "GB/T 7714-2015",
        "items": items,
    }
    REFS_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def validate_outputs(doc: Document, source_docx: Path) -> None:
    full_text = "\n".join(p.text for p in doc.paragraphs)
    payload = {
        "title": TITLE,
        "source_docx": str(source_docx),
        "main_md": str(MAIN_MD),
        "main_docx": str(MAIN_DOCX),
        "attachment_docx": str(ATTACHMENT_DOCX),
        "paragraph_count": len([p for p in doc.paragraphs if p.text.strip()]),
        "table_count": len(doc.tables),
        "image_count": sum(1 for rel in doc.part.rels.values() if "image" in rel.reltype),
        "has_cover": AUTHOR in full_text and SUPERVISOR in full_text and STUDENT_ID in full_text,
        "has_task_book": "毕业设计（论文）任务书" in full_text,
        "has_toc": "目 录" in full_text,
        "has_abstract": "摘  要" in full_text and "Abstract" in full_text,
        "has_conclusion": "结 论" in full_text,
        "header_text": HEADER_TEXT,
        "sensitive_hits": {
            token: full_text.count(token)
            for token in ["API_KEY", "POSTGRES_DSN", "sk-", "DASHSCOPE_API_KEY", "SILICONFLOW_API_KEY"]
            if token in full_text
        },
        "placeholder_hits": {
            token: full_text.count(token)
            for token in ["XXX", "YYY", "ZZZ", "[[TABLE:"]
            if token in full_text
        },
        "date": TODAY_CN,
    }
    CHECK_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    source_docx = locate_source_docx()
    shutil.copy2(source_docx, MAIN_DOCX)
    doc = Document(MAIN_DOCX)

    anchor = strip_old_front_matter(doc)
    insert_cover_and_taskbook(anchor)
    normalize_dates(doc)
    rename_chapter_titles(doc)
    rebuild_toc(doc)
    normalize_document_styles(doc)
    doc.save(MAIN_DOCX)

    build_attachment_docx()
    write_reference_json()
    validate_outputs(Document(MAIN_DOCX), source_docx)

    print(json.dumps({
        "main_docx": str(MAIN_DOCX),
        "attachment_docx": str(ATTACHMENT_DOCX),
        "reference_json": str(REFS_JSON),
        "check_json": str(CHECK_JSON),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
