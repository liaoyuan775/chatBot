from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path(r"D:/AgentLearning/chatBot/tmp/docs/frontmatter.docx")

TITLE = "智能语音聊天机器人设计与实现"
COLLEGE = "计算机科学与技术学院"
MAJOR = "计算机科学与技术"
CLASS_NAME = "________________"
AUTHOR = "廖园"
STUDENT_ID = "229074182"
SUPERVISOR = "王朋飞"
TODAY_CN = "2026年5月16日"


def set_run_font(run, east_asia: str = "宋体", ascii_font: str = "Times New Roman", size: float = 12, bold: bool = False) -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold


def format_para(paragraph, size: float = 12, center: bool = False, bold: bool = False, first_indent: bool = True, line: float = 18) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24) if (first_indent and not center) else Pt(0)
    for run in paragraph.runs:
        set_run_font(run, "黑体" if bold else "宋体", "Times New Roman", size, bold)


def add_line(doc: Document, text: str, size: float, center: bool = True, bold: bool = False, first_indent: bool = False, line: float = 24) -> None:
    p = doc.add_paragraph()
    p.add_run(text)
    format_para(p, size=size, center=center, bold=bold, first_indent=first_indent, line=line)


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.8)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    add_line(doc, "安徽工业大学", 22, center=True, bold=True)
    add_line(doc, "毕业设计（论文）说明书", 22, center=True, bold=True)
    add_line(doc, "", 12, center=True, bold=False)
    add_line(doc, TITLE, 18, center=True, bold=True)
    add_line(doc, "", 12, center=True, bold=False)
    add_line(doc, f"学院：{COLLEGE}", 14, center=True, bold=False)
    add_line(doc, f"专业：{MAJOR}", 14, center=True, bold=False)
    add_line(doc, f"班级：{CLASS_NAME}", 14, center=True, bold=False)
    add_line(doc, f"姓名：{AUTHOR}", 14, center=True, bold=False)
    add_line(doc, f"学号：{STUDENT_ID}", 14, center=True, bold=False)
    add_line(doc, f"指导教师：{SUPERVISOR}", 14, center=True, bold=False)
    add_line(doc, f"日期：{TODAY_CN}", 14, center=True, bold=False)
    add_page_break(doc)

    add_line(doc, "毕业设计（论文）任务书", 18, center=True, bold=True, line=18)
    for text in [
        f"题目：{TITLE}",
        "主要任务：围绕智能语音聊天机器人系统完成需求分析、数据库设计、总体架构设计、详细模块设计、功能实现与系统测试，形成符合学校模板要求的毕业设计论文。",
        "研究重点：重点说明统一会话运行时 SessionRuntime、integrated-realtime 与 split-chain 双语音策略、WebSocket 实时通话、RAG 检索增强、上下文记忆、Provider 与模型管理、人格配置和音色管理等关键设计。",
        "技术路线：前端采用 React、Vite 与 TypeScript 构建界面与运行时；后端采用 FastAPI、异步 SQLAlchemy、PostgreSQL 与 pgvector 构建会话服务、知识增强和实时通信能力。",
        "完成要求：论文需包含中英文摘要、目录、需求分析、数据库设计、总体设计、详细设计与实现、系统测试、参考文献、致谢，并配套系统架构图、功能模块图、语音流程图、实时通话时序图、数据库 E-R 图和真实页面截图。",
    ]:
        p = doc.add_paragraph()
        p.add_run(text)
        format_para(p, size=12, center=False, bold=False, first_indent=True, line=18)

    doc.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    main()
