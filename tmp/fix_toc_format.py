from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


DOC_PATH = Path(r"D:\AgentLearning\chatBot\output\doc\智能语音聊天机器人设计与实现.docx")


def set_run_fonts(run, east_asia: str, ascii_font: str = "Times New Roman", size_pt: float | None = None, bold=None):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def set_style_font(style, east_asia: str, size_pt: float, bold=None):
    style.font.name = "Times New Roman"
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), east_asia)
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold


def clear_tab_stops(paragraph_format):
    ppr = paragraph_format._element
    tabs = ppr.find(qn("w:tabs"))
    if tabs is not None:
        ppr.remove(tabs)


def add_right_dot_tab(paragraph_format, pos_cm: float):
    paragraph_format.tab_stops.add_tab_stop(Cm(pos_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


doc = Document(str(DOC_PATH))

# 1. 目录标题样式
toc_title = doc.paragraphs[34]
toc_title.style = doc.styles["Title"]
toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
toc_title.paragraph_format.space_before = Pt(24)
toc_title.paragraph_format.space_after = Pt(18)
toc_title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
for run in toc_title.runs:
    set_run_fonts(run, east_asia="黑体", size_pt=18, bold=True)

# 2. 目录样式：小四宋体、固定18磅、缩进和点线右对齐位置
style_specs = {
    "toc 1": {"left_cm": 0.0, "tab_cm": 15.55},
    "toc 2": {"left_cm": 0.74, "tab_cm": 15.55},
    "toc 3": {"left_cm": 1.48, "tab_cm": 15.55},
}

for style_name, spec in style_specs.items():
    style = doc.styles[style_name]
    pf = style.paragraph_format
    pf.left_indent = Cm(spec["left_cm"])
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    clear_tab_stops(pf)
    add_right_dot_tab(pf, spec["tab_cm"])
    set_style_font(style, east_asia="宋体", size_pt=12, bold=False)

# 3. 参考文献、致谢、附录设为一级标题，确保进入目录
for idx in [646, 660, 665]:
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        p.style = doc.styles["Heading 1"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before = True
        for run in p.runs:
            set_run_fonts(run, east_asia="黑体", size_pt=18, bold=True)

doc.save(str(DOC_PATH))
print(DOC_PATH)
