from pathlib import Path

from docx import Document


ZH_P1 = (
    "本课题面向中文语音交互场景，研究智能语音聊天机器人在多轮对话中的系统设计与实现。"
    "针对传统聊天机器人存在交互方式单一、语音链路割裂、上下文记忆不足以及私有知识利用能力较弱等问题，"
    "论文以统一会话运行时为主线，探索文本对话、语音输入、实时语音通话与知识增强能力的一体化组织方法。"
    "该研究对于提升语音交互的连续性、可维护性和工程可控性具有现实意义。"
)

ZH_P2 = (
    "本系统以前端 React、Vite、TypeScript 和后端 FastAPI 为主要开发技术，采用 PostgreSQL 与 pgvector "
    "构建会话数据、记忆数据和知识库向量数据的存储基础；在架构上采用前后端协同的分层设计，在方法上实现了统一会话运行时、"
    "split-chain 与 integrated-realtime 双语音链路、基于检索门控的检索增强生成流程，以及短期记忆、中期摘要和长期向量记忆相结合的上下文管理机制。"
    "系统完成了文本问答、语音输入、实时 WebSocket 通话、知识库问答、人格与音色配置等功能，并通过功能测试与流程验证对关键模块进行了检验。"
    "结果表明，系统能够较稳定地完成文本与语音两类主要交互任务，知识增强、实时通话和异常处理功能基本达到设计目标。"
)

ZH_P3 = (
    "研究表明，围绕统一会话运行时组织多模态语音对话能力，能够在保证系统扩展性的同时提升交互连续性与工程实现效率。"
    "本文的成果可为智能陪练、课程问答、资料助理等中文语音应用提供实现参考。"
    "仍需进一步优化的方面主要包括复杂噪声环境下的识别鲁棒性以及长时对话场景下的一致性保持。"
)

ZH_KW = "关键词：智能语音聊天机器人；多模态交互；检索增强生成；实时语音交互；统一会话运行时"

EN_P1 = (
    "This thesis focuses on Chinese voice interaction scenarios and studies the design and implementation of an intelligent voice chatbot for multi-turn dialogue. "
    "To address the limitations of traditional chatbots in single interaction modality, fragmented speech pipelines, insufficient contextual memory, and weak use of private knowledge, "
    "the study takes a unified session runtime as the main line and explores an integrated organization approach for text dialogue, voice input, real-time voice conversation, and knowledge enhancement. "
    "The work is of practical value for improving interaction continuity, maintainability, and engineering controllability in voice-based systems."
)

EN_P2 = (
    "The system is developed with React, Vite, and TypeScript on the frontend and FastAPI on the backend, while PostgreSQL and pgvector are used to store session data, memory data, and vectorized knowledge-base data. "
    "In terms of architecture, the system adopts a layered design with frontend-backend collaboration. In terms of methods, it implements a unified session runtime, two speech pipelines named split-chain and integrated-realtime, "
    "a retrieval-gated retrieval-augmented generation workflow, and a context management mechanism that combines short-term memory, mid-term summaries, and long-term vector memory. "
    "The system realizes text question answering, voice input, real-time WebSocket conversation, knowledge-based question answering, persona configuration, and voice profile configuration, and its key modules are verified through functional tests and workflow validation. "
    "The results show that the system can stably support both text and voice interaction, and that the knowledge enhancement, real-time conversation, and exception-handling functions basically meet the design objectives."
)

EN_P3 = (
    "The study shows that organizing multimodal voice dialogue around a unified session runtime can improve interaction continuity and engineering efficiency while preserving system extensibility. "
    "The implementation can provide a practical reference for Chinese voice applications such as intelligent tutoring, course question answering, and document assistants. "
    "Further improvement is still needed in speech recognition robustness under complex noise conditions and consistency maintenance in long-session dialogue."
)

EN_KW = "Keywords: intelligent voice chatbot; multimodal interaction; retrieval-augmented generation; real-time voice interaction; unified session runtime"


def update_markdown(md_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8")
    start_zh = text.index("## 摘要")
    start_en = text.index("## Abstract")
    start_ch1 = text.index("## 第1章")
    new_zh = (
        "## 摘要\n\n"
        f"{ZH_P1}\n\n"
        f"{ZH_P2}\n\n"
        f"{ZH_P3}\n\n"
        f"**{ZH_KW}**\n\n"
    )
    new_en = (
        "## Abstract\n\n"
        f"{EN_P1}\n\n"
        f"{EN_P2}\n\n"
        f"{EN_P3}\n\n"
        f"**{EN_KW}**\n\n"
    )
    text = text[:start_zh] + new_zh + new_en + text[start_ch1:]
    md_path.write_text(text, encoding="utf-8")


def update_docx(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    replacements = {
        33: ZH_P1,
        34: ZH_P2,
        35: ZH_P3,
        37: ZH_KW,
        39: EN_P1,
        41: EN_P2,
        42: EN_P3,
        43: EN_KW,
    }
    for idx, value in replacements.items():
        p = doc.paragraphs[idx]
        if p.runs:
            p.runs[0].text = value
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.add_run(value)
    doc.save(str(docx_path))


base = Path(r"D:\AgentLearning\chatBot\output\doc")
md_path = base / "智能语音聊天机器人设计与实现.md"
docx_path = max([p for p in base.glob("*.docx") if "附件" not in p.name], key=lambda p: p.stat().st_size)

update_markdown(md_path)
update_docx(docx_path)

print(md_path)
print(docx_path)
