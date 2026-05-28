from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn


BASE = Path(r"D:\AgentLearning\chatBot\output\doc")
DOCX_PATH = max([p for p in BASE.glob("*.docx") if "附件" not in p.name], key=lambda p: p.stat().st_size)


def remove_paragraph(paragraph):
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def clear_runs_and_hyperlinks(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_plain_text(paragraph, text):
    clear_runs_and_hyperlinks(paragraph)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is not None:
        rfonts.set(qn("w:ascii"), "Times New Roman")
        rfonts.set(qn("w:hAnsi"), "Times New Roman")
        rfonts.set(qn("w:eastAsia"), "宋体")


doc = Document(str(DOCX_PATH))

# Remove citations whose references were deleted.
for idx in [195, 201, 205, 223]:
    p = doc.paragraphs[idx]
    cleaned = re.sub(r"\[(12|13|14|15|16)\]", "", p.text).rstrip()
    set_plain_text(p, cleaned)

# Keep only the first strict reference block [1]-[11].
# Current duplicate/invalid tail is [12]-[16] + duplicated [1]-[11].
for _ in range(16):  # remove original paragraphs 674..689, always deleting current 674
    remove_paragraph(doc.paragraphs[674])

doc.save(str(DOCX_PATH))
print(DOCX_PATH)
