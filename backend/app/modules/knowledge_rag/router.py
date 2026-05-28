from __future__ import annotations

import io
import uuid
from pathlib import Path

from docx import Document as DocxDocument
from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field
from pypdf import PdfReader
from sqlalchemy import delete, select, update
from sqlalchemy.ext.asyncio import AsyncSession
from pptx import Presentation

from app.core.config import get_settings
from app.core.database import get_db
from app.core.provider_gateway import ProviderError, embedding, rerank
from app.models import KnowledgeBaseEntity, KnowledgeChunkEntity, KnowledgeDocumentEntity, SystemSettingEntity
from app.schemas.common import RetrievalConfigUpsert

settings = get_settings()
router = APIRouter(prefix="/api/knowledge", tags=["knowledge"])
DEFAULT_RETRIEVAL_CONFIG = {
    "top_k": 3,
    "similarity_threshold": 0.0,
    "rag_timeout_ms": 1500,
    "no_result_message": "未查询到相关知识，我将为您进行通用解答。",
    "embedding_provider": "siliconflow",
    "embedding_model": settings.siliconflow_embedding_model,
    "rerank_provider": "siliconflow",
    "rerank_model": settings.siliconflow_rerank_model,
}
DEFAULT_BASE_NAME = "默认知识库"


class KnowledgeBaseUpsert(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    description: str | None = None
    is_default: bool = False


def split_chunks(text: str, chunk_size: int = 500, overlap: int = 80) -> list[str]:
    chunks: list[str] = []
    i = 0
    while i < len(text):
        end = min(len(text), i + chunk_size)
        piece = text[i:end]
        if piece.strip():
            chunks.append(piece)
        if end == len(text):
            break
        i = max(0, end - overlap)
    return chunks


def _extract_text_from_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


def _extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _extract_text_from_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])


def _extract_text_from_pptx(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text)
    return "\n".join(texts)


def extract_text(file_name: str, data: bytes) -> tuple[str, str]:
    suffix = Path(file_name).suffix.lower().replace(".", "")
    if suffix in {"txt", "md", "markdown"}:
        return "txt", _extract_text_from_txt(data)
    if suffix == "pdf":
        return "pdf", _extract_text_from_pdf(data)
    if suffix == "docx":
        return "docx", _extract_text_from_docx(data)
    if suffix == "pptx":
        return "pptx", _extract_text_from_pptx(data)
    raise HTTPException(status_code=400, detail="Unsupported format. Allowed: txt, md, pdf, docx, pptx")


async def get_retrieval_config(db: AsyncSession) -> dict:
    row = await db.get(SystemSettingEntity, "knowledge_retrieval_config")
    if not row:
        row = SystemSettingEntity(setting_key="knowledge_retrieval_config", setting_value=DEFAULT_RETRIEVAL_CONFIG)
        db.add(row)
        await db.commit()
        await db.refresh(row)
        return dict(DEFAULT_RETRIEVAL_CONFIG)
    merged = {**DEFAULT_RETRIEVAL_CONFIG, **(row.setting_value or {})}
    if merged != (row.setting_value or {}):
        row.setting_value = merged
        await db.commit()
    return merged


async def ensure_default_base(db: AsyncSession) -> KnowledgeBaseEntity:
    row = await db.scalar(select(KnowledgeBaseEntity).where(KnowledgeBaseEntity.is_default.is_(True)))
    if row:
        return row
    row = await db.scalar(select(KnowledgeBaseEntity).order_by(KnowledgeBaseEntity.created_at.asc()))
    if row:
        row.is_default = True
        await db.commit()
        return row
    row = KnowledgeBaseEntity(name=DEFAULT_BASE_NAME, description="系统默认知识库", is_default=True)
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return row


async def resolve_knowledge_base(db: AsyncSession, knowledge_base_id: str | None) -> KnowledgeBaseEntity:
    if not knowledge_base_id:
        return await ensure_default_base(db)
    try:
        kb_uuid = uuid.UUID(knowledge_base_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Invalid knowledge_base_id.") from exc
    base = await db.get(KnowledgeBaseEntity, kb_uuid)
    if not base:
        raise HTTPException(status_code=404, detail="Knowledge base not found.")
    return base


async def process_document(
    db: AsyncSession,
    document: KnowledgeDocumentEntity,
    file_name: str,
    data: bytes,
    chunk_size: int,
    overlap: int,
    embedding_provider: str,
    embedding_model: str,
) -> tuple[str, int]:
    file_type, full_text = extract_text(file_name, data)
    if not full_text.strip():
        raise HTTPException(status_code=400, detail="No extractable content.")
    await db.execute(delete(KnowledgeChunkEntity).where(KnowledgeChunkEntity.document_id == document.id))
    chunks = split_chunks(full_text, chunk_size=chunk_size, overlap=overlap)
    vectors = await embedding(db, embedding_provider, embedding_model, chunks)
    for idx, chunk in enumerate(chunks):
        vec = vectors[idx] if idx < len(vectors) else vectors[-1]
        db.add(
            KnowledgeChunkEntity(
                document_id=document.id,
                chunk_index=idx,
                content=chunk,
                metadata_json={"summary": chunk[:120]},
                embedding=vec,
            )
        )
    document.file_name = file_name
    document.file_type = file_type
    document.parse_status = "parsed"
    document.size_bytes = len(data)
    document.chunk_count = len(chunks)
    await db.commit()
    return file_type, len(chunks)


@router.get("/bases")
async def list_knowledge_bases(db: AsyncSession = Depends(get_db)):
    await ensure_default_base(db)
    rows = (await db.scalars(select(KnowledgeBaseEntity).order_by(KnowledgeBaseEntity.created_at.asc()))).all()
    return [
        {
            "id": str(item.id),
            "name": item.name,
            "description": item.description or "",
            "is_default": item.is_default,
            "created_at": item.created_at,
            "updated_at": item.updated_at,
        }
        for item in rows
    ]


@router.post("/bases")
async def create_knowledge_base(payload: KnowledgeBaseUpsert = Body(...), db: AsyncSession = Depends(get_db)):
    exists = await db.scalar(select(KnowledgeBaseEntity).where(KnowledgeBaseEntity.name == payload.name.strip()))
    if exists:
        raise HTTPException(status_code=400, detail="Knowledge base name already exists.")
    if payload.is_default:
        await db.execute(update(KnowledgeBaseEntity).values(is_default=False))
    row = KnowledgeBaseEntity(
        name=payload.name.strip(),
        description=(payload.description or "").strip() or None,
        is_default=payload.is_default,
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return {"id": str(row.id), "message": "知识库创建成功。"}


@router.patch("/bases/{knowledge_base_id}")
async def update_knowledge_base(
    knowledge_base_id: uuid.UUID,
    payload: KnowledgeBaseUpsert = Body(...),
    db: AsyncSession = Depends(get_db),
):
    row = await db.get(KnowledgeBaseEntity, knowledge_base_id)
    if not row:
        raise HTTPException(status_code=404, detail="Knowledge base not found.")
    if payload.name.strip() != row.name:
        exists = await db.scalar(select(KnowledgeBaseEntity).where(KnowledgeBaseEntity.name == payload.name.strip()))
        if exists:
            raise HTTPException(status_code=400, detail="Knowledge base name already exists.")
    if payload.is_default:
        await db.execute(update(KnowledgeBaseEntity).values(is_default=False))
    row.name = payload.name.strip()
    row.description = (payload.description or "").strip() or None
    row.is_default = payload.is_default
    await db.commit()
    return {"message": "知识库更新成功。"}


@router.delete("/bases/{knowledge_base_id}")
async def delete_knowledge_base(knowledge_base_id: uuid.UUID, db: AsyncSession = Depends(get_db)):
    row = await db.get(KnowledgeBaseEntity, knowledge_base_id)
    if not row:
        raise HTTPException(status_code=404, detail="Knowledge base not found.")
    fallback = await db.scalar(
        select(KnowledgeBaseEntity).where(KnowledgeBaseEntity.id != knowledge_base_id).order_by(KnowledgeBaseEntity.created_at.asc())
    )
    if not fallback:
        raise HTTPException(status_code=400, detail="At least one knowledge base is required.")
    await db.execute(
        update(KnowledgeDocumentEntity)
        .where(KnowledgeDocumentEntity.knowledge_base_id == knowledge_base_id)
        .values(knowledge_base_id=fallback.id)
    )
    if row.is_default:
        fallback.is_default = True
    await db.delete(row)
    await db.commit()
    return {"message": "知识库删除成功，文档已迁移到其它知识库。"}


@router.get("/documents")
async def list_documents(
    knowledge_base_id: uuid.UUID | None = Query(default=None),
    db: AsyncSession = Depends(get_db),
):
    query = select(KnowledgeDocumentEntity).order_by(KnowledgeDocumentEntity.created_at.desc())
    if knowledge_base_id:
        query = query.where(KnowledgeDocumentEntity.knowledge_base_id == knowledge_base_id)
    rows = (await db.scalars(query)).all()
    base_map = {
        str(item.id): item
        for item in (await db.scalars(select(KnowledgeBaseEntity))).all()
    }
    return [
        {
            "id": str(r.id),
            "knowledge_base_id": str(r.knowledge_base_id) if r.knowledge_base_id else None,
            "knowledge_base_name": base_map.get(str(r.knowledge_base_id)).name if r.knowledge_base_id and base_map.get(str(r.knowledge_base_id)) else "",
            "file_name": r.file_name,
            "file_type": r.file_type,
            "parse_status": r.parse_status,
            "size_bytes": r.size_bytes,
            "chunk_count": r.chunk_count,
            "created_at": r.created_at,
        }
        for r in rows
    ]


@router.post("/documents")
async def upload_document(
    file: UploadFile = File(...),
    chunk_size: int = Form(default=500),
    overlap: int = Form(default=80),
    knowledge_base_id: str | None = Form(default=None),
    db: AsyncSession = Depends(get_db),
):
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    target_base = await resolve_knowledge_base(db, knowledge_base_id)
    file_name = file.filename or f"doc-{uuid.uuid4()}.txt"
    file_type = Path(file_name).suffix.lower().replace(".", "") or "txt"

    old = await db.scalar(
        select(KnowledgeDocumentEntity).where(
            KnowledgeDocumentEntity.file_name == file_name,
            KnowledgeDocumentEntity.knowledge_base_id == target_base.id,
        )
    )
    if old:
        await db.execute(delete(KnowledgeChunkEntity).where(KnowledgeChunkEntity.document_id == old.id))
        await db.delete(old)
        await db.commit()

    doc = KnowledgeDocumentEntity(
        knowledge_base_id=target_base.id,
        file_name=file_name,
        file_type=file_type,
        parse_status="parsing",
        size_bytes=len(data),
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)
    retrieval_config = await get_retrieval_config(db)
    try:
        _, chunk_count = await process_document(
            db,
            doc,
            file_name,
            data,
            chunk_size,
            overlap,
            embedding_provider=str(retrieval_config.get("embedding_provider", "siliconflow")),
            embedding_model=str(retrieval_config.get("embedding_model", settings.siliconflow_embedding_model)),
        )
    except ProviderError as exc:
        doc.parse_status = "failed"
        await db.commit()
        raise HTTPException(status_code=502, detail=f"Embedding failed: {exc}") from exc
    except HTTPException:
        doc.parse_status = "failed"
        await db.commit()
        raise
    except Exception:
        doc.parse_status = "failed"
        await db.commit()
        raise
    return {"id": str(doc.id), "chunk_count": chunk_count, "message": "文档上传并解析成功。"}


@router.delete("/documents/{document_id}")
async def delete_document(document_id: uuid.UUID, db: AsyncSession = Depends(get_db)):
    doc = await db.get(KnowledgeDocumentEntity, document_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")
    await db.execute(delete(KnowledgeChunkEntity).where(KnowledgeChunkEntity.document_id == document_id))
    await db.delete(doc)
    await db.commit()
    return {"message": "文档删除成功。"}


@router.post("/documents/{document_id}/reparse")
async def reparse_document(
    document_id: uuid.UUID,
    chunk_size: int = Form(default=500),
    overlap: int = Form(default=80),
    db: AsyncSession = Depends(get_db),
):
    doc = await db.get(KnowledgeDocumentEntity, document_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")
    old_chunks = (
        await db.scalars(
            select(KnowledgeChunkEntity)
            .where(KnowledgeChunkEntity.document_id == document_id)
            .order_by(KnowledgeChunkEntity.chunk_index.asc())
        )
    ).all()
    if not old_chunks:
        raise HTTPException(status_code=400, detail="文档没有已解析的分块内容，请删除后重新上传原始文件。")
    combined = "\n".join(chunk.content for chunk in old_chunks)
    if not combined.strip():
        raise HTTPException(status_code=400, detail="文档分块内容为空，请删除后重新上传原始文件。")

    # Re-parse directly from extracted text — do NOT call extract_text()
    # which expects raw binary (PDF/DOCX headers) and would fail on plain text.
    retrieval_config = await get_retrieval_config(db)
    emb_provider = str(retrieval_config.get("embedding_provider", "siliconflow"))
    emb_model = str(retrieval_config.get("embedding_model", settings.siliconflow_embedding_model))

    new_chunks = split_chunks(combined, chunk_size=chunk_size, overlap=overlap)
    from app.core.provider_gateway import ProviderError, embedding
    vectors = await embedding(db, emb_provider, emb_model, new_chunks)

    await db.execute(
        delete(KnowledgeChunkEntity).where(KnowledgeChunkEntity.document_id == document_id)
    )
    for idx, chunk_text in enumerate(new_chunks):
        vec = vectors[idx] if idx < len(vectors) else vectors[-1]
        db.add(
            KnowledgeChunkEntity(
                document_id=document_id,
                chunk_index=idx,
                content=chunk_text,
                metadata_json={"summary": chunk_text[:120]},
                embedding=vec,
            )
        )
    doc.parse_status = "parsed"
    doc.chunk_count = len(new_chunks)
    await db.commit()
    return {"message": "文档重解析成功。", "chunk_count": len(new_chunks)}


@router.put("/global-switch")
async def update_knowledge_switch(enabled: bool, db: AsyncSession = Depends(get_db)):
    row = await db.get(SystemSettingEntity, "knowledge_global")
    if not row:
        row = SystemSettingEntity(setting_key="knowledge_global", setting_value={"enabled": enabled})
        db.add(row)
    else:
        row.setting_value = {"enabled": enabled}
    await db.commit()
    return {"message": "Knowledge switch updated.", "enabled": enabled}


@router.get("/global-switch")
async def get_knowledge_switch(db: AsyncSession = Depends(get_db)):
    row = await db.get(SystemSettingEntity, "knowledge_global")
    return {"enabled": bool(row.setting_value.get("enabled", True)) if row else True}


@router.get("/retrieval-config")
async def fetch_retrieval_config(db: AsyncSession = Depends(get_db)):
    return await get_retrieval_config(db)


@router.put("/retrieval-config")
async def update_retrieval_config(payload: RetrievalConfigUpsert, db: AsyncSession = Depends(get_db)):
    row = await db.get(SystemSettingEntity, "knowledge_retrieval_config")
    if not row:
        row = SystemSettingEntity(setting_key="knowledge_retrieval_config", setting_value=payload.model_dump())
        db.add(row)
    else:
        row.setting_value = payload.model_dump()
    await db.commit()
    return {"message": "Retrieval config updated.", "config": payload.model_dump()}


@router.post("/retrieval-test")
async def retrieval_test(
    question: str = Form(...),
    top_k: int = Form(default=0),
    similarity_threshold: float = Form(default=999.0),
    knowledge_base_id: str | None = Form(default=None),
    db: AsyncSession = Depends(get_db),
):
    retrieval_config = await get_retrieval_config(db)
    if top_k <= 0:
        top_k = int(retrieval_config.get("top_k", DEFAULT_RETRIEVAL_CONFIG["top_k"]))
    if similarity_threshold == 999.0:
        similarity_threshold = float(
            retrieval_config.get("similarity_threshold", DEFAULT_RETRIEVAL_CONFIG["similarity_threshold"])
        )
    emb_provider = str(retrieval_config.get("embedding_provider", "siliconflow"))
    emb_model = str(retrieval_config.get("embedding_model", settings.siliconflow_embedding_model))
    rr_provider = str(retrieval_config.get("rerank_provider", "siliconflow"))
    rr_model = str(retrieval_config.get("rerank_model", settings.siliconflow_rerank_model))
    query_vec = (await embedding(db, emb_provider, emb_model, question))[0]
    query = select(KnowledgeChunkEntity).join(
        KnowledgeDocumentEntity,
        KnowledgeChunkEntity.document_id == KnowledgeDocumentEntity.id,
    )
    if knowledge_base_id:
        try:
            kb_uuid = uuid.UUID(knowledge_base_id)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="Invalid knowledge_base_id.") from exc
        query = query.where(KnowledgeDocumentEntity.knowledge_base_id == kb_uuid)
    rows = (await db.scalars(query)).all()

    def dot(a: list[float], b: list[float]) -> float:
        return sum(x * y for x, y in zip(a, b))

    scored = []
    for row in rows:
        score = dot(query_vec, row.embedding)
        if score >= similarity_threshold:
            scored.append({"row": row, "score": score})
    scored.sort(key=lambda x: x["score"], reverse=True)
    candidate_texts = [x["row"].content for x in scored[: max(top_k * 3, top_k)]]

    reranked = []
    try:
        reranked = await rerank(db, rr_provider, rr_model, question, candidate_texts, top_n=top_k)
    except Exception:  # noqa: BLE001
        reranked = []

    if reranked:
        results = [
            {"content": candidate_texts[item.get("index", 0)], "score": item.get("relevance_score", item.get("score", 0.0))}
            for item in reranked
        ]
    else:
        results = [{"content": x["row"].content, "score": round(x["score"], 4)} for x in scored[:top_k]]
    if not results:
        return {
            "results": [],
            "count": 0,
            "message": retrieval_config.get("no_result_message", DEFAULT_RETRIEVAL_CONFIG["no_result_message"]),
        }
    return {"results": results, "count": len(results)}
